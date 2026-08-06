import openpyxl
import pytest

from leap_projection.data import Fundamentals
from leap_projection.stockanalysis_xlsx import (
    enrich_fundamentals,
    load_stockanalysis_xlsx,
)

# Mirrors the real stockanalysis.com export layout: row 1 = ["Date", "TTM",
# then descending period-end dates], each following row = one line item.
ANNUAL_DATES = ["TTM", "2025-12-31", "2024-12-31", "2023-12-31", "2022-12-31", "2021-12-31", "2020-12-31"]


def _write_workbook(path, pe_values, eps_values, revenue_values, ebitda_margin_values):
    wb = openpyxl.Workbook()
    income = wb.active
    income.title = "Income-Annual"
    income.append(["Date"] + ANNUAL_DATES)
    income.append(["EPS (Diluted)"] + eps_values)
    income.append(["Revenue"] + revenue_values)
    income.append(["EBITDA Margin"] + ebitda_margin_values)

    ratios = wb.create_sheet("Ratios-Annual")
    ratios.append(["Date"] + ANNUAL_DATES)
    ratios.append(["PE Ratio"] + pe_values)

    wb.save(path)


def _fundamentals(**overrides) -> Fundamentals:
    base = dict(
        symbol="TEST",
        current_price=100.0,
        trailing_eps=4.0,
        forward_eps=4.5,
        trailing_pe=22.0,
        forward_pe=19.0,
        peg_ratio=1.1,
        earnings_growth=0.10,
        revenue_growth=0.08,
        analyst_target_mean=None,
        analyst_target_high=None,
        analyst_target_low=None,
        sector="Technology",
        market_cap=1_000_000_000,
    )
    base.update(overrides)
    return Fundamentals(**base)


def test_load_parses_annual_series_ascending(tmp_path):
    path = tmp_path / "test-financials.xlsx"
    # TTM, then 2025..2020 descending in the sheet -> ascending 2020..2025 once parsed
    _write_workbook(
        path,
        pe_values=[23.0, 25.0, 30.0, 40.0, 28.0, 50.0, 88.0],
        eps_values=[3.2, 2.6, 2.0, 1.2, 1.0, 1.1, 0.6],
        revenue_values=[48000, 45000, 39000, 33700, 31600, 29700, 25000],
        ebitda_margin_values=[0.30, 0.30, 0.27, 0.22, 0.19, 0.22, 0.19],
    )

    financials = load_stockanalysis_xlsx(str(path))

    assert list(financials.pe_ratio.index.year) == [2020, 2021, 2022, 2023, 2024, 2025]
    assert financials.pe_ratio.iloc[-1] == 25.0  # 2025
    assert financials.eps_diluted.iloc[-1] == 2.6
    assert financials.revenue.iloc[0] == 25000  # 2020, oldest


def test_load_raises_on_missing_required_sheets(tmp_path):
    path = tmp_path / "not-a-stockanalysis-file.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)

    with pytest.raises(ValueError):
        load_stockanalysis_xlsx(str(path))


def test_enrich_fundamentals_uses_historical_pe_median(tmp_path):
    path = tmp_path / "test-financials.xlsx"
    # Ascending by year: 2020=88, 2021=50, 2022=28, 2023=40, 2024=30, 2025=25.
    # Last 5 years (2021-2025) sorted: 25, 28, 30, 40, 50 -> median 30.
    _write_workbook(
        path,
        pe_values=[23.0, 25.0, 30.0, 40.0, 28.0, 50.0, 88.0],
        eps_values=[3.2, 2.6, 2.0, 1.2, 1.0, 1.1, 0.6],
        revenue_values=[48000, 45000, 39000, 33700, 31600, 29700, 25000],
        ebitda_margin_values=[0.30, 0.30, 0.27, 0.22, 0.19, 0.22, 0.19],
    )
    fundamentals = _fundamentals()

    enriched = enrich_fundamentals(fundamentals, str(path))

    assert enriched.historical_pe_multiple == pytest.approx(30.0)
    # EPS CAGR from 0.6 (2020) to 2.6 (2025) over 5 years
    assert enriched.earnings_growth == pytest.approx((2.6 / 0.6) ** (1 / 5) - 1, abs=1e-6)
    assert enriched.forward_eps == pytest.approx(4.0 * (1 + enriched.earnings_growth))
    # unrelated fields pass through untouched
    assert enriched.symbol == fundamentals.symbol
    assert enriched.trailing_pe == fundamentals.trailing_pe


def test_enrich_fundamentals_falls_back_when_pe_history_too_thin(tmp_path):
    path = tmp_path / "thin-financials.xlsx"
    # Only 2 non-missing annual PE values -> below MIN_PE_SAMPLES (3)
    _write_workbook(
        path,
        pe_values=[23.0, None, None, None, None, 50.0, 88.0],
        eps_values=[3.2, 2.6, 2.0, 1.2, 1.0, 1.1, 0.6],
        revenue_values=[48000, 45000, 39000, 33700, 31600, 29700, 25000],
        ebitda_margin_values=[0.30, 0.30, 0.27, 0.22, 0.19, 0.22, 0.19],
    )
    fundamentals = _fundamentals()

    enriched = enrich_fundamentals(fundamentals, str(path))

    assert enriched.historical_pe_multiple == fundamentals.historical_pe_multiple  # still None
