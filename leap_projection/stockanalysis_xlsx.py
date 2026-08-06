"""Optional enrichment from a stockanalysis.com "Export to Excel" workbook.

stockanalysis.com's export produces a workbook with 12 sheets
(Income/Balance-Sheet/Cash-Flow/Ratios x Annual/Quarterly/TTM), each shaped
the same way: row 1 holds period headers ("Date", "TTM", then descending
fiscal period-end dates), and every row below is one line item labeled in
column A.

This export does **not** include analyst estimates or price targets -- it's
historical financial statements and ratios only. What it's actually useful
for here is a genuine multi-year historical P/E series (replacing the
single-point-in-time proxy used when only FMP data is available) and a
longer earnings/revenue history for a steadier growth-rate estimate.

Usage: download a ticker's financials export from stockanalysis.com, then
pass its path via `--stockanalysis-xlsx` (CLI) or `enrich_fundamentals()`
(library). It's a manual download, not a live API, so it enriches a single
`Fundamentals` snapshot rather than being fetched automatically per run.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from typing import Optional

import openpyxl
import pandas as pd

from .data import Fundamentals, _cagr, _clip_growth

REQUIRED_SHEETS = {"Income-Annual", "Ratios-Annual"}

PE_LOOKBACK_YEARS = 5
GROWTH_LOOKBACK_YEARS = 5
MIN_PE_SAMPLES = 3


@dataclass
class StockAnalysisFinancials:
    pe_ratio: pd.Series  # annual, ascending by period-end date
    eps_diluted: pd.Series
    revenue: pd.Series
    ebitda_margin: pd.Series


def _sheet_row_series(ws, row_label: str) -> pd.Series:
    """Extract one line-item row as a Series indexed by period-end date
    (ascending), skipping the 'Date'/'TTM' header columns."""
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    target_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == row_label:
            target_row = r
            break
    if target_row is None:
        return pd.Series(dtype=float)

    dates, values = [], []
    for col, header in enumerate(headers, start=1):
        if header in (None, "Date", "TTM"):
            continue
        value = ws.cell(row=target_row, column=col).value
        if value is None:
            continue
        try:
            period_date = pd.Timestamp(header)
        except (ValueError, TypeError):
            continue
        dates.append(period_date)
        values.append(value)

    return pd.Series(values, index=pd.DatetimeIndex(dates, name="date")).sort_index()


def load_stockanalysis_xlsx(path: str) -> StockAnalysisFinancials:
    """Parse a stockanalysis.com financials export into annual series."""
    wb = openpyxl.load_workbook(path, data_only=True)
    missing = REQUIRED_SHEETS - set(wb.sheetnames)
    if missing:
        raise ValueError(
            f"'{path}' doesn't look like a stockanalysis.com financials export "
            f"(missing sheet(s): {', '.join(sorted(missing))})."
        )

    income = wb["Income-Annual"]
    ratios = wb["Ratios-Annual"]
    return StockAnalysisFinancials(
        pe_ratio=_sheet_row_series(ratios, "PE Ratio"),
        eps_diluted=_sheet_row_series(income, "EPS (Diluted)"),
        revenue=_sheet_row_series(income, "Revenue"),
        ebitda_margin=_sheet_row_series(income, "EBITDA Margin"),
    )


def _historical_pe_multiple(pe_series: pd.Series, years: int = PE_LOOKBACK_YEARS) -> Optional[float]:
    """Median of the trailing `years` annual P/E ratios, positive values only
    (a loss-making year's P/E is meaningless as a reversion target)."""
    recent = pe_series.tail(years)
    positive = recent[recent > 0]
    if len(positive) < MIN_PE_SAMPLES:
        return None
    return float(positive.median())


def _series_cagr(series: pd.Series, years: int = GROWTH_LOOKBACK_YEARS) -> Optional[float]:
    recent = series.dropna().tail(years + 1)
    if len(recent) < 2:
        return None
    oldest, newest = float(recent.iloc[0]), float(recent.iloc[-1])
    n_years = len(recent) - 1
    return _clip_growth(_cagr(oldest, newest, n_years))


def enrich_fundamentals(fundamentals: Fundamentals, xlsx_path: str) -> Fundamentals:
    """Return a copy of `fundamentals` enriched with a stockanalysis.com
    financials export: a multi-year historical P/E reversion multiple, and
    longer-run EPS/revenue CAGR where available (falling back to the
    original FMP-derived values when the export doesn't have enough
    history for a given field).
    """
    financials = load_stockanalysis_xlsx(xlsx_path)

    historical_pe = _historical_pe_multiple(financials.pe_ratio)
    eps_growth = _series_cagr(financials.eps_diluted)
    revenue_growth = _series_cagr(financials.revenue)

    earnings_growth = eps_growth if eps_growth is not None else fundamentals.earnings_growth
    revenue_growth = revenue_growth if revenue_growth is not None else fundamentals.revenue_growth

    forward_eps = (
        fundamentals.trailing_eps * (1 + earnings_growth)
        if fundamentals.trailing_eps and earnings_growth is not None
        else fundamentals.forward_eps
    )
    peg_ratio = (
        fundamentals.trailing_pe / (earnings_growth * 100)
        if fundamentals.trailing_pe and earnings_growth and earnings_growth > 0
        else fundamentals.peg_ratio
    )

    return replace(
        fundamentals,
        earnings_growth=earnings_growth,
        revenue_growth=revenue_growth,
        forward_eps=forward_eps,
        peg_ratio=peg_ratio,
        historical_pe_multiple=(
            historical_pe if historical_pe is not None else fundamentals.historical_pe_multiple
        ),
    )
